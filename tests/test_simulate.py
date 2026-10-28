import itertools

import numpy as np
import pandas as pd
import pytest

from football_model import simulate, spread

CONFERENCES = ["ACC", "Big 12", "Big Ten", "SEC", "Sun Belt"]


def make_league() -> tuple[dict[str, str], dict[str, str], pd.DataFrame]:
    """20 teams in 5 conferences of 4 (Sun Belt split into 2 divisions) + 1 independent.

    The schedule is a conference round robin (6 games per conference), one week each.
    """
    conferences, divisions, rows = {}, {}, []
    game_id = 1
    for c in CONFERENCES:
        members = [f"{c} {i}" for i in range(4)]
        for t in members:
            conferences[t] = c
        if c == "Sun Belt":
            divisions.update({members[0]: "East", members[1]: "East"})
            divisions.update({members[2]: "West", members[3]: "West"})
        for week, (h, a) in enumerate(itertools.combinations(members, 2), start=1):
            rows.append(
                {
                    "game_id": game_id,
                    "season": 2026,
                    "week": week,
                    "season_type": "regular",
                    "start_date": pd.Timestamp("2026-09-05", tz="UTC") + pd.Timedelta(weeks=week),
                    "home_team": h,
                    "away_team": a,
                    "home_conference": c,
                    "away_conference": c,
                    "conference_game": True,
                    "neutral_site": False,
                    "notes": None,
                }
            )
            game_id += 1
    conferences["Notre Dame"] = simulate.INDEPENDENTS
    return conferences, divisions, pd.DataFrame(rows)


def run(completed, remaining, conferences, divisions, n_sims=400, seed=1, mean=None):
    """Simulate with known ratings: team i of each conference is rated 10 - 5i."""
    teams = sorted(conferences) + [simulate.FCS_TEAM]
    if mean is None:
        values = [3.0] + [10.0 - 5 * int(t[-1]) if t[-1].isdigit() else 5.0 for t in teams]
        mean = pd.Series(values, index=[simulate.HOME_FIELD, *teams])
    ucov = np.eye(len(mean)) * 0.5
    return simulate.simulate_season(
        completed, remaining, mean, ucov, 14.0, conferences, divisions, n_sims, seed
    )


def split_schedule(schedule: pd.DataFrame, played_weeks: int):
    """Completed games (with margins: home team always wins by 7) and remaining games."""
    completed = schedule[schedule["week"] <= played_weeks].assign(margin=7)
    remaining = schedule[schedule["week"] > played_weeks]
    return completed, remaining


def test_probabilities_sum_correctly():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)

    odds = run(completed, remaining, conferences, divisions)

    by_conf = odds[odds["conference"] != simulate.INDEPENDENTS].groupby("conference")
    np.testing.assert_allclose(by_conf["p_conf_champ"].sum(), 1.0)
    np.testing.assert_allclose(by_conf["p_title_game"].sum(), 2.0)
    assert odds["p_playoff"].sum() == pytest.approx(simulate.PLAYOFF_TEAMS)


def test_power_champions_always_make_the_playoff():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)

    odds = run(completed, remaining, conferences, divisions)

    power = odds[odds["conference"].isin(simulate.POWER_CONFERENCES)]
    assert (power["p_playoff"] >= power["p_conf_champ"] - 1e-12).all()


def test_expected_wins_within_bounds():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)

    odds = run(completed, remaining, conferences, divisions).set_index("team")

    # Each team plays 3 games; at most one more as a title game.
    assert (odds["exp_regular_wins"] >= odds["wins"]).all()
    assert (odds["exp_regular_wins"] <= 3).all()
    assert (odds["exp_wins"] <= 4).all()


def test_same_seed_same_results():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 2)

    first = run(completed, remaining, conferences, divisions, seed=7)
    second = run(completed, remaining, conferences, divisions, seed=7)

    pd.testing.assert_frame_equal(first, second)


def test_completed_title_game_fixes_the_champion():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 6)
    title = completed.iloc[[0]].assign(
        game_id=999,
        week=15,
        home_team="SEC 3",
        away_team="SEC 0",
        margin=3,
        notes="SEC Championship",
        start_date=pd.Timestamp("2026-12-05", tz="UTC"),
    )
    completed = pd.concat([completed, title], ignore_index=True)

    odds = run(completed, remaining, conferences, divisions).set_index("team")

    # SEC 3 is the weakest SEC team but won the real title game.
    assert odds.loc["SEC 3", "p_conf_champ"] == 1.0
    assert odds.loc["SEC 0", "p_title_game"] == 1.0


def test_order_conference_head_to_head_breaks_two_way_tie():
    pct = np.array([0.75, 0.75, 0.5])
    rating = np.array([10.0, 0.0, 5.0])  # team 0 is rated higher
    beat = np.zeros((3, 3), dtype=int)
    beat[1, 0], beat[0, 1] = 1, -1  # but team 1 beat team 0

    order = simulate.order_conference(pct, rating, beat)

    assert list(order) == [1, 0, 2]


def test_order_conference_three_way_tie_uses_rating():
    pct = np.array([0.5, 0.5, 0.5])
    rating = np.array([1.0, 3.0, 2.0])
    beat = np.zeros((3, 3), dtype=int)
    beat[0, 1], beat[1, 0] = 1, -1

    assert list(simulate.order_conference(pct, rating, beat)) == [1, 2, 0]


def test_ridge_posterior_matches_fit_ratings_and_shrinks_with_alpha():
    games = pd.DataFrame(
        {
            "home_team": ["A", "B", "C"],
            "away_team": ["B", "C", "A"],
            "neutral_site": [False, False, True],
            "margin": [10, 3, -7],
            "start_date": pd.Timestamp("2026-09-01", tz="UTC"),
        }
    )
    prior = pd.Series([2.0, 1.0, 0.0, -1.0], index=[simulate.HOME_FIELD, "A", "B", "C"])

    mean, _ = simulate.ridge_posterior(games, prior, alpha=3.0)
    fit = spread.fit_ratings(games, prior, spread.SpreadParams(alpha=3.0))
    pd.testing.assert_series_equal(mean, fit)

    mean_big, ucov_big = simulate.ridge_posterior(games, prior, alpha=1e6)
    np.testing.assert_allclose(mean_big, prior, atol=1e-4)
    # (X'X + alpha I)^-1 ~ I / alpha when alpha dwarfs X'X: the prior dominates.
    np.testing.assert_allclose(ucov_big, np.eye(4) / 1e6, rtol=1e-3, atol=1e-10)


def test_split_noise_removes_rating_uncertainty():
    ucov = np.eye(3) * 0.1
    contrasts = np.array([[1.0, 1.0, -1.0]])  # c'Uc = 0.3

    sigma_e = simulate.split_noise(16.0, ucov, contrasts)

    assert sigma_e == pytest.approx(16.0 / np.sqrt(1.3))
    assert sigma_e < 16.0


def test_is_title_game_uses_notes_and_ignores_army_navy():
    games = pd.DataFrame(
        {
            "game_id": [1, 2, 3],
            "season": [2024, 2024, 2019],
            "week": [15, 16, 15],
            "start_date": pd.to_datetime(["2024-12-07", "2024-12-14", "2019-12-07"], utc=True),
            "home_team": ["Army", "Army", "LSU"],
            "away_team": ["Tulane", "Navy", "Georgia"],
            "home_conference": ["American Athletic", "American Athletic", "SEC"],
            "away_conference": ["American Athletic", "American Athletic", "SEC"],
            "conference_game": [True, True, True],
            "notes": ["American Athletic Championship", None, None],
        }
    )

    # Labeled title game; Army-Navy (lone December conference game, 2024) is not one;
    # before 2022 a lone December conference game is inferred to be the title game.
    assert list(simulate.is_title_game(games)) == [True, False, True]


def test_nine_wins_and_bowl_odds_are_consistent():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)

    odds = run(completed, remaining, conferences, divisions)

    # Teams play at most 4 games here, so nobody can reach 9 or 6 wins.
    assert (odds["p_9_wins"] == 0).all()
    assert (odds["p_bowl"] == 0).all()
    assert (odds["p_9_wins"] <= odds["p_bowl"]).all()


def test_odds_tabs_layout():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)
    odds = run(completed, remaining, conferences, divisions)

    tabs = simulate.odds_tabs(odds)

    assert list(tabs) == ["Playoff odds", *CONFERENCES, simulate.INDEPENDENTS]
    playoff = tabs["Playoff odds"]
    assert list(playoff["Rank"]) == list(range(1, len(odds) + 1))
    assert playoff["Playoff"].is_monotonic_decreasing
    sec = tabs["SEC"]
    assert "Conference" not in sec.columns
    assert sec["Win conference"].is_monotonic_decreasing
    assert {"Expected wins", "Win 9+", "Bowl eligible (6+)"} <= set(sec.columns)
    assert "Win conference" not in tabs[simulate.INDEPENDENTS].columns


def test_write_odds_excel_highlights_expected_wins(tmp_path):
    from openpyxl import load_workbook

    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)
    odds = run(completed, remaining, conferences, divisions)
    path = tmp_path / "odds.xlsx"

    simulate.write_odds_excel(odds, path)

    sheet = load_workbook(path)["SEC"]
    headers = [c.value for c in sheet[1]]
    wins_col = headers.index("Expected wins") + 1
    assert sheet.cell(2, wins_col).fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet.cell(2, headers.index("Playoff") + 1).number_format == "0%"


def test_same_odds_detects_reproduced_snapshot():
    conferences, divisions, schedule = make_league()
    completed, remaining = split_schedule(schedule, 3)
    odds = run(completed, remaining, conferences, divisions)
    old_snapshot = odds.drop(columns="p_9_wins")  # saved before the column existed

    assert simulate.same_odds(old_snapshot, odds)
    assert not simulate.same_odds(old_snapshot.assign(p_playoff=0.5), odds)
